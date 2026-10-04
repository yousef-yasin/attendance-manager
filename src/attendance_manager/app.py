"""Application entry point."""

from __future__ import annotations

import logging
import sys
import traceback
from types import TracebackType

from attendance_manager import APP_ID, APP_NAME, ORGANIZATION_NAME, __version__
from attendance_manager.i18n import tr
from attendance_manager.logging_setup import configure_logging
from attendance_manager.paths import AppPaths, resource_path

log = logging.getLogger("attendance_manager")


def _set_windows_app_id() -> None:
    """Give the taskbar our own icon/grouping instead of python.exe's."""
    if sys.platform == "win32":
        try:
            import ctypes

            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(f"{APP_ID}.{__version__}")
        except Exception:
            pass


def self_test(argv: list[str]) -> int:
    """Verify a (frozen) build end-to-end without touching real user data.

    ``AttendanceManager.exe --self-test [result-file]`` builds the UI off-screen,
    creates a course from generated Excel files, records a session, writes a
    report and a backup in a temporary folder, and exits 0 on success.
    """
    import os
    import shutil
    import tempfile
    from pathlib import Path

    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    result_file = next(
        (Path(a) for a in argv[argv.index("--self-test") + 1 :] if not a.startswith("-")), None
    )
    work = Path(tempfile.mkdtemp(prefix="am-selftest-"))
    lines: list[str] = []
    status = 1
    try:
        import openpyxl
        from PySide6.QtCharts import QChart  # noqa: F401  (bundled?)
        from PySide6.QtWidgets import QApplication

        from attendance_manager.importing.importer import load_file
        from attendance_manager.services.context import AppContext
        from attendance_manager.services.course_service import CourseDetails
        from attendance_manager.ui.main_window import MainWindow
        from attendance_manager.ui.theme import apply_theme

        master, weekly = work / "00.xlsx", work / "11.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["#", "Student Id", "Student Name"])
        for i in range(1, 6):
            wb.active.append([i, f"20240000{i}", f"Student {i} Test"])
        wb.save(master)
        wb = openpyxl.Workbook()
        wb.active.append(["Name", "Email"])
        for i in (1, 2, 3):
            wb.active.append([f"Student {i} Test", f"20240000{i}@uni.edu"])
        wb.save(weekly)

        app = QApplication.instance() or QApplication([argv[0]])
        apply_theme(app)  # type: ignore[arg-type]
        ctx = AppContext(AppPaths(work / "data"))
        loaded = load_file(master)
        roster = ctx.courses.roster_preview(loaded, loaded.detection.mapping)  # type: ignore[arg-type]
        course = ctx.courses.create_course(CourseDetails("Self test"), roster, master)
        loaded = load_file(weekly)
        preview = ctx.attendance.build_preview(course.id, loaded, loaded.detection.mapping)  # type: ignore[arg-type]
        import datetime as dt

        session = ctx.attendance.save_session(
            preview, title="Lecture 1", session_date=dt.date.today()
        )
        assert (session.present_count, session.absent_count) == (3, 2), "matching failed"
        window = MainWindow(ctx)
        window.show()
        for key in ("dashboard", "students", "sessions", "reports", "backup"):
            window.navigate(key)
            app.processEvents()
        report = ctx.reports.generate(course.id, work / "report.xlsx")
        assert openpyxl.load_workbook(report).sheetnames[0] == "Summary"
        backup = ctx.backups.create_backup(work / "backup.zip")
        ctx.backups.validate_backup(backup)
        window.close()
        ctx.close()
        lines.append(f"OK {APP_NAME} {__version__}: import, matching, UI, report and backup work")
        status = 0
    except Exception:
        lines.append("FAILED\n" + traceback.format_exc())
    finally:
        shutil.rmtree(work, ignore_errors=True)
    text = "\n".join(lines)
    if result_file is not None:
        result_file.write_text(text, encoding="utf-8")
    if sys.stdout is not None:
        print(text)
    return status


def main(argv: list[str] | None = None) -> int:
    from PySide6.QtCore import QLockFile, Qt, QTimer
    from PySide6.QtGui import QGuiApplication, QIcon
    from PySide6.QtWidgets import QApplication, QMessageBox

    from attendance_manager.persistence.database import SchemaTooNewError
    from attendance_manager.services.context import AppContext
    from attendance_manager.ui.theme import apply_theme

    argv = list(sys.argv if argv is None else argv)
    if "--self-test" in argv:
        return self_test(argv)
    paths = AppPaths.default().ensure()
    log_file = configure_logging(paths.logs_dir)
    log.info("Starting %s %s", APP_NAME, __version__)
    _set_windows_app_id()

    QGuiApplication.setHighDpiScaleFactorRoundingPolicy(
        Qt.HighDpiScaleFactorRoundingPolicy.PassThrough
    )
    app = QApplication(argv)
    app.setApplicationName(APP_NAME)
    app.setApplicationDisplayName(APP_NAME)
    app.setOrganizationName(ORGANIZATION_NAME)
    app.setApplicationVersion(__version__)
    icon_file = resource_path("ui/assets/app_icon.png")
    if icon_file.exists():
        app.setWindowIcon(QIcon(str(icon_file)))
    apply_theme(app)

    def excepthook(
        exc_type: type[BaseException], exc: BaseException, tb: TracebackType | None
    ) -> None:
        log.critical("Unhandled error:\n%s", "".join(traceback.format_exception(exc_type, exc, tb)))
        QMessageBox.critical(
            None,
            APP_NAME,
            tr("errors.crash", log=log_file),
        )

    sys.excepthook = excepthook

    # One instance at a time, so two windows never write to the same database.
    lock = QLockFile(str(paths.root / "app.lock"))
    lock.setStaleLockTime(0)
    if not lock.tryLock(100):
        QMessageBox.information(None, APP_NAME, tr("app.already_open"))
        return 0

    try:
        ctx = AppContext(paths)
    except SchemaTooNewError:
        QMessageBox.critical(None, APP_NAME, tr("errors.schema_too_new"))
        return 1
    except Exception as exc:
        log.exception("Could not open the database")
        QMessageBox.critical(
            None,
            APP_NAME,
            tr("errors.database_open", error=exc, backups=paths.backups_dir),
        )
        return 1

    from attendance_manager.ui.language import apply_language, load_language
    from attendance_manager.ui.main_window import MainWindow

    # Language is applied before any widget exists, so the first paint is already
    # in the saved language and direction.
    apply_language(app, load_language(ctx))
    window = MainWindow(ctx)
    window.show()

    def startup_backup() -> None:
        try:
            if ctx.backups.auto_backup_if_due():
                log.info("Automatic backup created")
        except Exception:
            log.exception("Automatic backup failed")

    QTimer.singleShot(2500, startup_backup)
    try:
        return app.exec()
    finally:
        ctx.close()
        lock.unlock()
        log.info("Exited")


if __name__ == "__main__":
    raise SystemExit(main())
