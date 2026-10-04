"""Low-level spreadsheet reading.

The format is detected from the file *content*, not the extension, because
exported "xls" files are frequently really xlsx, HTML or CSV/TSV documents.
"""

from __future__ import annotations

import contextlib
import csv
import datetime as dt
import io
import logging
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from attendance_manager.domain.errors import ImportFileError
from attendance_manager.i18n import tr

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".xls", ".xlsx", ".xlsm", ".csv", ".tsv", ".txt")
MAX_FILE_BYTES = 50 * 1024 * 1024
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


@dataclass
class SheetData:
    name: str
    rows: list[list[Any]]

    @property
    def n_rows(self) -> int:
        return len(self.rows)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def cell(self, row: int, col: int) -> Any:
        if row < len(self.rows) and col < len(self.rows[row]):
            return self.rows[row][col]
        return None


@dataclass
class WorkbookData:
    path: Path
    file_format: str
    sheets: list[SheetData]


def read_workbook(path: Path | str) -> WorkbookData:
    path = Path(path)
    if not path.exists():
        raise ImportFileError(tr("errors.import.not_found", file=path.name))
    if not path.is_file():
        raise ImportFileError(tr("errors.import.not_a_file", file=path.name))
    size = path.stat().st_size
    if size == 0:
        raise ImportFileError(tr("errors.import.empty", file=path.name))
    if size > MAX_FILE_BYTES:
        raise ImportFileError(tr("errors.import.too_large", file=path.name))
    try:
        head = path.read_bytes()[:2048]
    except OSError as exc:
        raise ImportFileError(
            tr("errors.import.locked", file=path.name),
            details=str(exc),
        ) from exc

    if head.startswith(_OLE2_MAGIC):
        sheets, fmt = _read_xls(path), "xls"
    elif head.startswith(b"PK"):
        sheets, fmt = _read_xlsx(path), "xlsx"
    else:
        text = _decode_text(path.read_bytes())
        if text.lstrip()[:1] == "<" and "<t" in text.lower():
            sheets, fmt = _read_html(text), "html"
        else:
            sheets, fmt = _read_delimited(text), "csv"

    sheets = [s for s in (_trim(sheet) for sheet in sheets) if s.rows]
    if not sheets:
        raise ImportFileError(tr("errors.import.no_data", file=path.name))
    log.info("Read %s (%s): %s", path.name, fmt, [(s.name, s.n_rows, s.n_cols) for s in sheets])
    return WorkbookData(path=path, file_format=fmt, sheets=sheets)


def _read_xls(path: Path) -> list[SheetData]:
    import xlrd

    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except xlrd.biffh.XLRDError as exc:
        message = str(exc)
        if "encrypt" in message.lower():
            raise ImportFileError(tr("errors.import.password", file=path.name)) from exc
        raise ImportFileError(
            tr("errors.import.invalid_xls", file=path.name), details=message
        ) from exc
    except Exception as exc:
        raise ImportFileError(
            tr("errors.import.invalid_xls", file=path.name), details=str(exc)
        ) from exc

    sheets: list[SheetData] = []
    try:
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            rows: list[list[Any]] = []
            for r in range(sheet.nrows):
                row: list[Any] = []
                for c in range(sheet.ncols):
                    ctype = sheet.cell_type(r, c)
                    value = sheet.cell_value(r, c)
                    if ctype == xlrd.XL_CELL_DATE:
                        with contextlib.suppress(Exception):
                            value = xlrd.xldate_as_datetime(value, book.datemode)
                    elif ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        value = None
                    elif ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    row.append(value)
                rows.append(row)
            sheets.append(SheetData(name=sheet.name, rows=rows))
    finally:
        book.release_resources()
    return sheets


def _read_xlsx(path: Path) -> list[SheetData]:
    import openpyxl

    # Load from memory: openpyxl rejects real .xlsx content saved with a ".xls" name.
    data = io.BytesIO(path.read_bytes())
    try:
        book = openpyxl.load_workbook(data, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError(
            tr("errors.import.invalid_xlsx", file=path.name), details=str(exc)
        ) from exc
    try:
        sheets = []
        for ws in book.worksheets:
            rows = [list(row) for row in ws.iter_rows(values_only=True)]
            sheets.append(SheetData(name=ws.title, rows=rows))
        return sheets
    finally:
        book.close()


def _decode_text(data: bytes) -> str:
    for bom, encoding in (
        (b"\xff\xfe", "utf-16"),
        (b"\xfe\xff", "utf-16"),
        (b"\xef\xbb\xbf", "utf-8-sig"),
    ):
        if data.startswith(bom):
            return data.decode(encoding, errors="replace")
    for encoding in ("utf-8", "cp1256", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _read_delimited(text: str) -> list[SheetData]:
    if "\x00" in text[:1000]:
        raise ImportFileError(tr("errors.import.binary"))
    sample = text[:10000]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel_tab if sample.count("\t") > sample.count(",") else csv.excel
    rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    return [SheetData(name="Sheet1", rows=rows)]


class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append("".join(self._cell).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None and self.tables:
            self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def _read_html(text: str) -> list[SheetData]:
    parser = _TableParser()
    parser.feed(text)
    return [
        SheetData(name=f"Table {i + 1}", rows=[list(r) for r in table])
        for i, table in enumerate(parser.tables)
        if table
    ]


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _trim(sheet: SheetData) -> SheetData:
    """Drop trailing empty rows/columns (Excel often reports a huge used range)."""
    rows = [list(r) for r in sheet.rows]
    while rows and all(_is_blank(v) for v in rows[-1]):
        rows.pop()
    width = 0
    for row in rows:
        for idx in range(len(row) - 1, -1, -1):
            if not _is_blank(row[idx]):
                width = max(width, idx + 1)
                break
    rows = [(r + [None] * width)[:width] for r in rows]
    return SheetData(name=sheet.name, rows=rows)


def coerce_datetime(value: Any) -> dt.datetime | None:
    if isinstance(value, dt.datetime):
        return value
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time())
    return None
