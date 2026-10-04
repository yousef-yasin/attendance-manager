"""Professional multi-sheet Excel attendance report."""

from __future__ import annotations

import datetime as dt
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from attendance_manager import APP_NAME, __version__
from attendance_manager.domain.models import (
    AttendanceStatus,
    Course,
    SessionParticipant,
    Student,
    StudentStanding,
)
from attendance_manager.domain.statistics import CourseStatistics, StudentStats, Thresholds
from attendance_manager.i18n import format_date, is_rtl, tr

SHEET_NAMES = (
    "Summary",
    "Cumulative Attendance",
    "Session by Session",
    "Absence Ranking",
    "Highest Attendance",
    "At-Risk Students",
    "Student Details",
)

INK = "1B2437"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=18, color=INK)
SUBTITLE_FONT = Font(name="Calibri", size=11, color="52514E")
SECTION_FONT = Font(name="Calibri", bold=True, size=13, color=INK)
LABEL_FONT = Font(name="Calibri", bold=True, color="52514E")
THIN = Side(style="thin", color="D9DCE3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRESENT_FILL = PatternFill("solid", fgColor="DDF3E4")
ABSENT_FILL = PatternFill("solid", fgColor="FBE0E0")
STANDING_FILLS = {
    StudentStanding.PERFECT: PatternFill("solid", fgColor="D5F0DC"),
    StudentStanding.GOOD: PatternFill("solid", fgColor="E6F0FB"),
    StudentStanding.WARNING: PatternFill("solid", fgColor="FFF1CC"),
    StudentStanding.AT_RISK: PatternFill("solid", fgColor="FAD9D9"),
    StudentStanding.NO_SESSIONS: PatternFill("solid", fgColor="EEEEEE"),
}
PCT_FORMAT = "0.0%"


@dataclass(frozen=True)
class ReportData:
    course: Course
    statistics: CourseStatistics
    info_columns: Sequence[str]
    thresholds: Thresholds
    instructor: str = ""
    generated_at: dt.datetime | None = None


SHEET_IDS = (
    "summary",
    "cumulative",
    "by_session",
    "absence_ranking",
    "highest",
    "at_risk",
    "details",
)


def sheet_name(index: int) -> str:
    """Localized worksheet name (Excel allows at most 31 characters, no []:*?/\\)."""
    name = tr(f"report.sheet.{SHEET_IDS[index]}")
    for ch in "[]:*?/\\":
        name = name.replace(ch, " ")
    return name[:31]


def _standing(s: StudentStats) -> str:
    return tr(f"standing.{s.standing.name.lower()}")


def _pct(value: float | None) -> float | None:
    return None if value is None else round(value / 100.0, 6)


def _write_header(ws: Worksheet, row: int, headers: Sequence[str], start_col: int = 1) -> None:
    for offset, text in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + offset, value=text)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def _write_rows(
    ws: Worksheet,
    start_row: int,
    rows: Sequence[Sequence[Any]],
    pct_cols: Sequence[int] = (),
    start_col: int = 1,
) -> int:
    r = start_row
    for values in rows:
        for offset, value in enumerate(values):
            cell = ws.cell(row=r, column=start_col + offset, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")
            if offset in pct_cols:
                cell.number_format = PCT_FORMAT
        r += 1
    return r


def _autosize(ws: Worksheet, min_width: int = 8, max_width: int = 48) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = max(len(line) for line in str(cell.value).split("\n"))
            if isinstance(cell.value, float):
                length = 7
            widths[cell.column] = max(widths.get(cell.column, 0), length)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(
            min_width, min(max_width, width + 2)
        )


def _table(ws: Worksheet, header_row: int, last_row: int, n_cols: int) -> None:
    if last_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{last_row}"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)


def _student_cols(student: Student, info_columns: Sequence[str]) -> list[Any]:
    return [
        student.student_number,
        student.full_name,
        *[student.info.get(c, "") for c in info_columns],
    ]


def _stat_cols(s: StudentStats) -> list[Any]:
    return [s.sessions, s.present, s.absent, _pct(s.percentage), _standing(s)]


def write_course_report(path: Path, data: ReportData) -> Path:
    stats = data.statistics
    students = [s for s in stats.students if s.student.is_active or s.sessions > 0]
    wb = Workbook()
    summary = wb.active
    assert summary is not None
    summary.title = sheet_name(0)
    _summary_sheet(summary, data, students)

    base_headers = [tr("report.col.id"), tr("report.col.name"), *data.info_columns]
    stat_headers = [
        tr("report.col.sessions"),
        tr("report.col.present"),
        tr("report.col.absent"),
        tr("report.col.percent"),
        tr("report.col.status"),
    ]
    n_base = len(base_headers)

    # 2. cumulative
    ws = wb.create_sheet(sheet_name(1))
    headers = [*base_headers, *stat_headers]
    _write_header(ws, 1, headers)
    rows = [[*_student_cols(s.student, data.info_columns), *_stat_cols(s)] for s in students]
    last = _write_rows(ws, 2, rows, pct_cols=[n_base + 3]) - 1
    _color_status(ws, 2, last, len(headers), students)
    _table(ws, 1, last, len(headers))
    _autosize(ws)

    # 3. session by session
    ws = wb.create_sheet(sheet_name(2))
    sessions = stats.sessions
    session_headers = [
        f"#{s.session.session_number} {s.session.label}\n{format_date(s.session.session_date)}"
        for s in sessions
    ]
    headers = [
        tr("report.col.id"),
        tr("report.col.name"),
        *session_headers,
        tr("report.col.present"),
        tr("report.col.absent"),
        tr("report.col.percent"),
    ]
    _write_header(ws, 1, headers)
    r = 2
    for s in students:
        row: list[Any] = [s.student.student_number, s.student.full_name]
        for sess in sessions:
            status = s.history.get(sess.session.id)
            row.append(
                tr("report.code_present")
                if status is AttendanceStatus.PRESENT
                else tr("report.code_absent")
                if status is AttendanceStatus.ABSENT
                else "—"
            )
        row += [s.present, s.absent, _pct(s.percentage)]
        _write_rows(ws, r, [row], pct_cols=[len(row) - 1])
        for i, sess in enumerate(sessions):
            cell = ws.cell(row=r, column=3 + i)
            cell.alignment = Alignment(horizontal="center")
            status = s.history.get(sess.session.id)
            if status is AttendanceStatus.PRESENT:
                cell.fill = PRESENT_FILL
            elif status is AttendanceStatus.ABSENT:
                cell.fill = ABSENT_FILL
        r += 1
    last_student_row = r - 1
    totals: list[tuple[str, list[Any], bool]] = [
        (tr("report.col.present"), [s.present for s in sessions], False),
        (tr("report.col.absent"), [s.absent for s in sessions], False),
        (tr("report.col.percent"), [_pct(s.percentage) for s in sessions], True),
    ]
    for label, values, is_pct in totals:
        ws.cell(row=r, column=2, value=label).font = LABEL_FONT
        for i, value in enumerate(values):
            cell = ws.cell(row=r, column=3 + i, value=value)
            cell.font = Font(bold=True)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
            if is_pct:
                cell.number_format = PCT_FORMAT
        r += 1
    ws.cell(row=r + 1, column=1, value=tr("report.grid_legend")).font = SUBTITLE_FONT
    if last_student_row >= 2:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_student_row}"
    ws.freeze_panes = "C2"
    _autosize(ws, max_width=36)
    for i in range(len(sessions)):
        ws.column_dimensions[get_column_letter(3 + i)].width = 16

    # 4. absence ranking
    ws = wb.create_sheet(sheet_name(3))
    ranked = sorted(
        (s for s in students if s.sessions), key=lambda s: (-s.absent, s.percentage or 0.0)
    )
    _ranking_sheet(
        ws,
        ranked,
        [
            tr("report.col.rank"),
            tr("report.col.id"),
            tr("report.col.name"),
            tr("report.col.absent"),
            tr("report.col.present"),
            tr("report.col.sessions"),
            tr("report.col.percent"),
            tr("report.col.status"),
        ],
        lambda s: [s.absent, s.present, s.sessions, _pct(s.percentage), _standing(s)],
        rank_value=lambda s: s.absent,
    )

    # 5. highest attendance
    ws = wb.create_sheet(sheet_name(4))
    ranked = sorted(
        (s for s in students if s.sessions), key=lambda s: (-(s.percentage or 0.0), -s.present)
    )
    _ranking_sheet(
        ws,
        ranked,
        [
            tr("report.col.rank"),
            tr("report.col.id"),
            tr("report.col.name"),
            tr("report.col.percent"),
            tr("report.col.present"),
            tr("report.col.absent"),
            tr("report.col.sessions"),
            tr("report.col.status"),
        ],
        lambda s: [_pct(s.percentage), s.present, s.absent, s.sessions, _standing(s)],
        rank_value=lambda s: round(s.percentage or 0.0, 6),
    )

    # 6. at-risk
    ws = wb.create_sheet(sheet_name(5))
    ws["A1"] = tr("report.sheet.at_risk")
    ws["A1"].font = TITLE_FONT
    ws["A2"] = tr(
        "report.at_risk_note",
        at_risk=f"{data.thresholds.at_risk_below:g}",
        warning=f"{data.thresholds.warning_below:g}",
    )
    ws["A2"].font = SUBTITLE_FONT
    at_risk = stats.dashboard.at_risk_students + stats.dashboard.warning_students
    headers = [
        *base_headers,
        tr("report.col.absent"),
        tr("report.col.present"),
        tr("report.col.sessions"),
        tr("report.col.percent"),
        tr("report.col.status"),
    ]
    _write_header(ws, 4, headers)
    if at_risk:
        rows = [
            [
                *_student_cols(s.student, data.info_columns),
                s.absent,
                s.present,
                s.sessions,
                _pct(s.percentage),
                _standing(s),
            ]
            for s in at_risk
        ]
        last = _write_rows(ws, 5, rows, pct_cols=[n_base + 3]) - 1
        _color_status(ws, 5, last, len(headers), at_risk)
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last}"
    else:
        ws.cell(row=5, column=1, value=tr("report.no_at_risk")).font = SUBTITLE_FONT
    ws.freeze_panes = "C5"
    _autosize(ws)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 14)

    # 7. student details (long format)
    ws = wb.create_sheet(sheet_name(6))
    headers = [
        *base_headers,
        tr("report.col.session_number"),
        tr("report.col.session"),
        tr("report.col.date"),
        tr("report.col.attendance"),
        tr("report.col.overall_percent"),
        tr("report.col.overall_status"),
    ]
    _write_header(ws, 1, headers)
    r = 2
    for s in students:
        for sess in sessions:
            status = s.history.get(sess.session.id)
            if status is None:
                continue
            row = [
                *_student_cols(s.student, data.info_columns),
                sess.session.session_number,
                sess.session.label,
                sess.session.session_date,
                tr("status.present") if status is AttendanceStatus.PRESENT else tr("status.absent"),
                _pct(s.percentage),
                _standing(s),
            ]
            _write_rows(ws, r, [row], pct_cols=[len(row) - 2])
            ws.cell(row=r, column=n_base + 3).number_format = "dd mmm yyyy"
            ws.cell(row=r, column=n_base + 4).fill = (
                PRESENT_FILL if status is AttendanceStatus.PRESENT else ABSENT_FILL
            )
            r += 1
    _table(ws, 1, r - 1, len(headers))
    _autosize(ws)
    ws.column_dimensions[get_column_letter(n_base + 3)].width = 14

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = sheet.title == sheet_name(2)
        sheet.sheet_view.rightToLeft = is_rtl()
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToHeight = 0
    wb.properties.creator = f"{APP_NAME} {__version__}"
    wb.properties.title = tr("report.properties_title", course=data.course.display_name)
    return _save(wb, path)


def _save(wb: Workbook, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f"~{path.stem}.tmp.xlsx")
    wb.save(temp)
    temp.replace(path)
    return path


def _color_status(
    ws: Worksheet, first: int, last: int, status_col: int, students: Sequence[StudentStats]
) -> None:
    for i, s in enumerate(students):
        row = first + i
        if row > last:
            break
        ws.cell(row=row, column=status_col).fill = STANDING_FILLS[s.standing]


def _ranking_sheet(
    ws: Worksheet,
    ranked: Sequence[StudentStats],
    headers: Sequence[str],
    values: Callable[[StudentStats], list[Any]],
    rank_value: Callable[[StudentStats], object],
) -> None:
    _write_header(ws, 1, headers)
    rows = []
    rank = 0
    previous: object = object()
    for index, s in enumerate(ranked, start=1):
        value = rank_value(s)
        if value != previous:
            rank, previous = index, value
        rows.append([rank, s.student.student_number, s.student.full_name, *values(s)])
    pct_col = list(headers).index(tr("report.col.percent"))
    last = _write_rows(ws, 2, rows, pct_cols=[pct_col]) - 1
    _color_status(ws, 2, last, len(headers), ranked)
    _table(ws, 1, last, len(headers))
    _autosize(ws)
    ws.column_dimensions["A"].width = 7


def _summary_sheet(ws: Worksheet, data: ReportData, students: Sequence[StudentStats]) -> None:
    stats = data.statistics
    d = stats.dashboard
    generated = data.generated_at or dt.datetime.now()
    ws["A1"] = tr("report.title")
    ws["A1"].font = TITLE_FONT
    ws["A2"] = data.course.display_name
    ws["A2"].font = Font(name="Calibri", size=13, bold=True, color="1F3A5F")
    ws["A3"] = tr("report.generated", when=format_date(generated, "long_datetime")) + (
        f" · {data.instructor}" if data.instructor else ""
    )
    ws["A3"].font = SUBTITLE_FONT

    ws["A5"] = tr("report.course")
    ws["A5"].font = SECTION_FONT
    course_rows = [
        (tr("report.course_name"), data.course.name),
        (tr("course_dialog.section"), data.course.section or "—"),
        (tr("course_dialog.semester"), data.course.semester or "—"),
        (tr("course_dialog.year"), data.course.year or "—"),
    ]
    r = 6
    for label, value in course_rows:
        ws.cell(row=r, column=1, value=label).font = LABEL_FONT
        ws.cell(row=r, column=2, value=value)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value=tr("report.key_figures")).font = SECTION_FONT
    r += 1
    highest_names = tr("common.list_separator").join(
        s.student.full_name for s in d.highest_absence_students[:5]
    )
    kpis: list[tuple[str, Any, str | None]] = [
        (tr("report.kpi.students"), d.total_students, None),
        (tr("report.kpi.sessions"), d.total_sessions, None),
        (tr("report.kpi.average"), _pct(d.average_percentage), PCT_FORMAT),
        (tr("report.kpi.present"), d.total_present, None),
        (tr("report.kpi.absent"), d.total_absent, None),
        (tr("report.kpi.perfect"), len(d.perfect_students), None),
        (tr("report.kpi.with_absences"), len(d.students_with_absences), None),
        (tr("report.kpi.highest"), _pct(d.highest_percentage), PCT_FORMAT),
        (tr("report.kpi.highest_absences"), d.highest_absence_count, None),
        (tr("report.kpi.most_absent_names"), highest_names or "—", None),
        (
            tr("report.kpi.at_risk", limit=f"{data.thresholds.at_risk_below:g}"),
            len(d.at_risk_students),
            None,
        ),
        (
            tr("report.kpi.warning", limit=f"{data.thresholds.warning_below:g}"),
            len(d.warning_students),
            None,
        ),
    ]
    for label, value, fmt in kpis:
        ws.cell(row=r, column=1, value=label).font = LABEL_FONT
        cell = ws.cell(row=r, column=2, value=value if value is not None else "—")
        cell.alignment = Alignment(horizontal="right" if is_rtl() else "left")
        if fmt and value is not None:
            cell.number_format = fmt
        r += 1

    r += 1
    ws.cell(row=r, column=1, value=tr("report.sessions")).font = SECTION_FONT
    r += 1
    header_row = r
    _write_header(
        ws,
        r,
        [
            "#",
            tr("report.col.session"),
            tr("report.col.date"),
            tr("report.col.file"),
            tr("report.col.present"),
            tr("report.col.absent"),
            tr("report.col.percent"),
        ],
    )
    rows = [
        [
            s.session.session_number,
            s.session.label,
            s.session.session_date,
            s.session.original_filename,
            s.present,
            s.absent,
            _pct(s.percentage),
        ]
        for s in stats.sessions
    ]
    last = _write_rows(ws, r + 1, rows, pct_cols=[6]) - 1
    for row in range(r + 1, last + 1):
        ws.cell(row=row, column=3).number_format = "dd mmm yyyy"

    if stats.sessions:
        chart = LineChart()
        chart.title = tr("report.chart_trend")
        chart.y_axis.title = tr("report.col.percent")
        chart.y_axis.number_format = "0%"
        chart.y_axis.scaling.min = 0
        chart.y_axis.scaling.max = 1
        chart.height, chart.width = 7.5, 16
        chart.legend = None
        values = Reference(ws, min_col=7, min_row=header_row, max_row=last)
        cats = Reference(ws, min_col=2, min_row=header_row + 1, max_row=last)
        chart.add_data(values, titles_from_data=True)
        chart.set_categories(cats)
        series = chart.series[0]
        series.graphicalProperties.line.solidFill = "2A78D6"
        series.graphicalProperties.line.width = 28000
        series.marker.symbol = "circle"
        series.marker.size = 7
        series.marker.graphicalProperties.solidFill = "2A78D6"
        series.marker.graphicalProperties.line.solidFill = "FFFFFF"
        ws.add_chart(chart, "E5")

    dist_row = last + 3
    ws.cell(row=dist_row, column=1, value=tr("report.distribution")).font = SECTION_FONT
    _write_header(ws, dist_row + 1, [tr("report.col.range"), tr("report.col.students")])
    dist_rows = [
        [tr(f"distribution.bucket{i}"), count] for i, (_, count) in enumerate(d.distribution)
    ]
    last_dist = _write_rows(ws, dist_row + 2, dist_rows) - 1
    if students:
        bar = BarChart()
        bar.type = "col"
        bar.title = tr("report.chart_distribution")
        bar.legend = None
        bar.height, bar.width = 7.5, 16
        bar.add_data(
            Reference(ws, min_col=2, min_row=dist_row + 1, max_row=last_dist), titles_from_data=True
        )
        bar.set_categories(Reference(ws, min_col=1, min_row=dist_row + 2, max_row=last_dist))
        bar.series[0].graphicalProperties.solidFill = "2A78D6"
        bar.series[0].graphicalProperties.line.noFill = True
        bar.gapWidth = 60
        ws.add_chart(bar, f"E{dist_row}")

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 22
    for col in "EFG":
        ws.column_dimensions[col].width = 13


def write_session_report(
    path: Path,
    course: Course,
    session_label: str,
    session_date: dt.date,
    original_filename: str,
    present: Sequence[Student],
    absent: Sequence[Student],
    unmatched: Sequence[SessionParticipant],
) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = tr("report.session_sheet")
    ws.sheet_view.rightToLeft = is_rtl()
    ws["A1"] = session_label
    ws["A1"].font = TITLE_FONT
    ws["A2"] = tr(
        "report.session_subtitle",
        course=course.display_name,
        date=format_date(session_date, "long"),
        file=original_filename,
    )
    ws["A2"].font = SUBTITLE_FONT
    total = len(present) + len(absent)
    figures = [
        (tr("report.col.present"), len(present)),
        (tr("report.col.absent"), len(absent)),
        (tr("report.col.percent"), (len(present) / total) if total else None),
        (tr("report.not_on_list"), len(unmatched)),
    ]
    for i, (label, value) in enumerate(figures):
        ws.cell(row=4 + i, column=1, value=label).font = LABEL_FONT
        cell = ws.cell(row=4 + i, column=2, value=value)
        if label == tr("report.col.percent"):
            cell.number_format = PCT_FORMAT
    r = 9
    _write_header(ws, r, [tr("report.col.id"), tr("report.col.name"), tr("report.col.status")])
    present_text, absent_text = tr("status.present"), tr("status.absent")
    rows = [[s.student_number, s.full_name, present_text] for s in present]
    rows += [[s.student_number, s.full_name, absent_text] for s in absent]
    last = _write_rows(ws, r + 1, rows) - 1
    for row in range(r + 1, last + 1):
        ws.cell(row=row, column=3).fill = (
            PRESENT_FILL if ws.cell(row=row, column=3).value == present_text else ABSENT_FILL
        )
    if unmatched:
        u = wb.create_sheet(tr("report.not_on_list")[:31])
        u.sheet_view.rightToLeft = is_rtl()
        _write_header(
            u,
            1,
            [tr("attendance.col.row"), tr("attendance.col.identifier"), tr("attendance.col.name")],
        )
        _write_rows(u, 2, [[p.row_number, p.raw_identifier, p.name] for p in unmatched])
        _autosize(u)
    _autosize(ws)
    ws.column_dimensions["A"].width = 20
    return _save(wb, path)
