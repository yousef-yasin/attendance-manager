"""Localization: catalogs, completeness, no hard-coded UI text, runtime switching, RTL, persistence."""

from __future__ import annotations

import ast
import datetime as dt
import os
import re
import string
from pathlib import Path

import openpyxl
import pytest

from attendance_manager.i18n import (
    LANGUAGES,
    format_date,
    load_catalog,
    plural_category,
    set_language,
    strip_bidi,
    tr,
    trn,
)
from attendance_manager.reporting.excel_report import SHEET_IDS, SHEET_NAMES

from .conftest import SYNTHETIC, add_session, create_course

SRC = Path(__file__).resolve().parent.parent / "src" / "attendance_manager"
PLURAL_FORMS = re.compile(r"\.(zero|one|two|few|many|other)$")
ARABIC = re.compile(r"[؀-ۿ]")


@pytest.fixture(autouse=True)
def english_afterwards():
    yield
    set_language("en")


def _base(key: str) -> str:
    return PLURAL_FORMS.sub("", key)


def _placeholders(text: str) -> set[str]:
    return {f[1] for f in string.Formatter().parse(text) if f[1]}


def _keys_used_in_code() -> set[str]:
    keys: set[str] = set()
    for path in SRC.rglob("*.py"):
        keys |= set(
            re.findall(r'\b(?:tr|trn|tr_list)\(\s*"([^"]+)"', path.read_text(encoding="utf-8"))
        )
    return keys


# ------------------------------------------------------------------ catalogs
def test_english_translations_exist():
    en = load_catalog("en")
    assert len(en) > 400
    assert tr("nav.dashboard") == "Dashboard"
    assert tr("actions.create_course") == "Create course"


def test_arabic_translations_exist_and_use_required_terminology():
    set_language("ar")
    expected = {
        "nav.dashboard": "لوحة التحكم",
        "nav.courses": "المواد الدراسية",
        "nav.students": "الطلاب",
        "nav.reports": "التقارير",
        "nav.settings": "الإعدادات",
        "nav.backup": "النسخ الاحتياطي والاستعادة",
        "actions.create_course": "إنشاء مادة",
        "actions.add_attendance": "إضافة حضور",
        "dashboard.kpi.students": "إجمالي الطلاب",
        "status.present": "حاضر",
        "status.absent": "غائب",
        "attendance.card.rate": "نسبة الحضور",
        "dashboard.kpi.highest": "أعلى نسبة حضور",
        "dashboard.kpi.most_absences": "أعلى عدد من الغيابات",
        "standing.at_risk": "يحتاج إلى متابعة",
        "actions.save": "حفظ",
        "actions.cancel": "إلغاء",
        "actions.close": "إغلاق",
        "students.filter_tooltip": "تصفية",
        "course_dialog.step2": "2  قائمة الطلاب الأساسية",
    }
    for key, text in expected.items():
        assert tr(key) == text, key
    assert tr("reports.generate").startswith("تصدير التقرير")


def test_no_missing_translation_keys():
    en = {_base(k) for k in load_catalog("en")}
    ar = {_base(k) for k in load_catalog("ar")}
    assert en == ar, f"only en: {sorted(en - ar)}; only ar: {sorted(ar - en)}"
    missing = sorted(k for k in _keys_used_in_code() if k not in en)
    assert not missing, f"keys used in code but missing from the catalogs: {missing}"


def test_placeholders_match_between_languages():
    en, ar = load_catalog("en"), load_catalog("ar")
    for key, text in en.items():
        if not isinstance(text, str):
            continue
        base = _base(key)
        ar_forms = [v for k, v in ar.items() if _base(k) == base]
        for ar_text in ar_forms:
            extra = _placeholders(ar_text) - _placeholders(text) - {"count"}
            assert not extra, f"{key}: Arabic uses unknown placeholders {extra}"
        if not PLURAL_FORMS.search(key):
            assert _placeholders(ar[key]) == _placeholders(text), key


def test_every_arabic_string_is_translated():
    """Arabic values must contain Arabic script, except format-only / brand strings."""
    allowed = {
        "nav.tooltip",
        "common.list_separator",
        "dashboard.list.risk_row",
        "course_dialog.section_placeholder",
    }
    ar = load_catalog("ar")
    untranslated = [
        k
        for k, v in ar.items()
        if isinstance(v, str)
        and re.search(
            r"[A-Za-z]{3}",
            re.sub(
                r"\{[^}]*\}|<[^>]+>|Excel|Ctrl|\.[a-z]+|\*\.\w+|USB|Microsoft Teams|CSV|[\w.]+@[\w.]+|\d+",
                "",
                v,
            ),
        )
        and not ARABIC.search(v)
        and k not in allowed
        and not k.startswith(("distribution.", "date.format", "date.qt_format"))
    ]
    assert not untranslated, untranslated


def test_report_sheet_names_match_catalog():
    set_language("en")
    assert tuple(tr(f"report.sheet.{i}") for i in SHEET_IDS) == SHEET_NAMES
    set_language("ar")
    for i in SHEET_IDS:
        name = tr(f"report.sheet.{i}")
        assert ARABIC.search(name) and len(name) <= 31


# -------------------------------------------------------------------- plurals
@pytest.mark.parametrize(
    ("count", "category"),
    [
        (0, "zero"),
        (1, "one"),
        (2, "two"),
        (3, "few"),
        (10, "few"),
        (11, "many"),
        (99, "many"),
        (100, "other"),
        (102, "other"),
        (103, "few"),
        (111, "many"),
    ],
)
def test_arabic_plural_rules(count, category):
    assert plural_category("ar", count) == category


def test_plural_forms():
    assert trn("plural.students", 1) == "1 student"
    assert trn("plural.students", 45) == "45 students"
    set_language("ar")
    assert trn("plural.students", 1) == "طالب واحد"
    assert trn("plural.students", 2) == "طالبان"
    assert trn("plural.students", 5) == "5 طلاب"
    assert trn("plural.students", 45) == "45 طالبًا"
    assert trn("plural.students", 100) == "100 طالب"


def test_dates_are_localized_with_western_digits():
    day = dt.date(2026, 7, 18)
    assert format_date(day, "weekday_date") == "Sat 18 Jul 2026"
    set_language("ar")
    assert format_date(day, "weekday_date") == "السبت 18 يوليو 2026"
    assert format_date(day, "long") == "السبت، 18 يوليو 2026"


def test_data_is_never_translated():
    """Course names/sections/IDs are data; only the surrounding labels change."""
    from attendance_manager.domain.models import Course, CourseStatus

    now = dt.datetime(2026, 7, 1, tzinfo=dt.timezone.utc)
    course = Course(
        1, "u", "مهارات حياتية", "1", "Summer", "2026", "", CourseStatus.ACTIVE, now, now
    )
    assert course.display_name == "مهارات حياتية · Section 1 · Summer 2026"
    set_language("ar")
    assert course.display_name == "مهارات حياتية · الشعبة 1 · Summer 2026"


# ------------------------------------------------------- no hard-coded UI text
# call name -> positional argument indexes that are shown to the user
USER_FACING_ARGS: dict[str, tuple[int, ...]] = {
    "label": (0,),
    "button": (0, 4),
    "QLabel": (0,),
    "QPushButton": (0,),
    "QCheckBox": (0,),
    "setText": (0,),
    "setWindowTitle": (0,),
    "setPlaceholderText": (0,),
    "setToolTip": (0,),
    "addItem": (0,),
    "addTab": (1,),
    "addAction": (1,),
    "Banner": (1,),
    "StatCard": (0,),
    "EmptyState": (1, 2),
    "add_header": (0, 1),
    "Column": (0,),
    "toast": (1, 2),
    "show_error": (2,),
    "confirm": (1, 2, 3),
    "show_busy": (0,),
    "setSpecialValueText": (0,),
    "setSuffix": (0,),
    "information": (2,),
    "critical": (2,),
    "warning": (2,),
    "getText": (1, 2),
    "getSaveFileName": (1, 3),
    "getOpenFileName": (1, 3),
    "DropZone": (0, 1, 2),
    "ListRow": (0, 1),
    "StatusChip": (1,),
    "set_value": (1,),
    "QBarSet": (0,),
    "setFormat": (0,),
    "ValidationError": (0,),
    "ImportFileError": (0,),
    "BackupError": (0,),
    "AttendanceError": (0,),
    "PreviewNotice": (1,),
}
USER_FACING_KWARGS = {"text", "title", "message", "details", "hint", "tooltip"}


def _literal_text(node: ast.AST) -> str | None:
    if isinstance(node, ast.Constant) and isinstance(node.value, str):
        return node.value
    if isinstance(node, ast.JoinedStr):
        return "".join(
            v.value for v in node.values if isinstance(v, ast.Constant) and isinstance(v.value, str)
        )
    return None


def test_no_hardcoded_user_facing_strings():
    offenders = []
    targets = [
        *SRC.joinpath("ui").rglob("*.py"),
        *SRC.joinpath("services").rglob("*.py"),
        *SRC.joinpath("importing").rglob("*.py"),
        SRC / "app.py",
    ]
    for path in targets:
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            func = node.func
            name = func.attr if isinstance(func, ast.Attribute) else getattr(func, "id", "")
            if name not in USER_FACING_ARGS:
                continue
            candidates = [node.args[i] for i in USER_FACING_ARGS[name] if i < len(node.args)]
            candidates += [kw.value for kw in node.keywords if kw.arg in USER_FACING_KWARGS]
            for arg in candidates:
                text = _literal_text(arg)
                # Machine codes such as "duplicate-course" are not prose.
                if (
                    text
                    and re.search(r"[A-Za-z]{2,}", text)
                    and not re.fullmatch(r"[a-z0-9_.-]+", text)
                ):
                    offenders.append(f"{path.relative_to(SRC)}:{node.lineno}: {name}({text!r})")
    assert not offenders, "Hard-coded user-facing text:\n" + "\n".join(offenders)


# ------------------------------------------------------------ Qt integration
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


@pytest.fixture
def window(qtbot, qapp, ctx):
    from attendance_manager.ui.language import apply_language, load_language
    from attendance_manager.ui.main_window import MainWindow
    from attendance_manager.ui.theme import apply_theme

    apply_theme(qapp)
    apply_language(qapp, load_language(ctx))
    course = create_course(ctx, SYNTHETIC / "00.xls")
    add_session(ctx, course.id, SYNTHETIC / "11.xls")
    win = MainWindow(ctx)
    qtbot.addWidget(win)
    win.show()
    yield win
    apply_language(qapp, "en")


def _all_texts(widget) -> list[str]:
    from PySide6.QtWidgets import QAbstractButton, QLabel

    texts = [w.text() for w in widget.findChildren(QLabel) if w.text()]
    texts += [w.text() for w in widget.findChildren(QAbstractButton) if w.text()]
    return [strip_bidi(t) for t in texts]  # ignore invisible direction marks


def test_language_switches_at_runtime_and_rtl_is_applied(window, qapp):
    from PySide6.QtCore import Qt

    assert qapp.layoutDirection() == Qt.LayoutDirection.LeftToRight
    assert window.nav_buttons["dashboard"].text().strip() == "Dashboard"
    window.change_language("ar")
    assert qapp.layoutDirection() == Qt.LayoutDirection.RightToLeft
    assert window.isRightToLeft()
    assert window.nav_buttons["dashboard"].text().strip() == "لوحة التحكم"
    assert window.page_title.text() == "لوحة التحكم"
    for key in ("courses", "sessions", "students", "reports", "settings", "backup"):
        window.navigate(key)
        assert ARABIC.search(window.page_title.text()), key
    # Every visible label/button on every page is Arabic, except data (numbers, IDs, names, files).
    window.navigate("dashboard")
    english_left = [
        t
        for t in _all_texts(window)
        if re.search(r"[A-Za-z]{3,}", t)
        and not ARABIC.search(t)
        and not re.search(
            r"\.(xls|xlsx|csv)|Course|[A-Z]{4,}|English|Adam|Basma|Carim|Dana|Emad|Farah|"
            r"Ghaith|Hala|Iyad|Jana|Khaled|Layan|Summer|Lecture|Ctrl|"
            r"^[A-Z]:\\|[0-9a-f]{8}-[0-9a-f]{4}-",  # data: folder paths, installation id
            t,
        )
    ]
    assert not english_left, english_left


def test_ltr_and_english_are_restored(window, qapp):
    from PySide6.QtCore import Qt

    window.change_language("ar")
    window.change_language("en")
    assert qapp.layoutDirection() == Qt.LayoutDirection.LeftToRight
    assert not window.isRightToLeft()
    assert window.nav_buttons["sessions"].text().strip() == "Attendance Sessions"
    assert not [t for t in _all_texts(window) if ARABIC.search(t) and t != LANGUAGES["ar"]]


def test_state_is_kept_when_switching(window):
    window.navigate("students")
    course_id = window.state.course_id
    window.change_language("ar")
    assert window.current_page_key() == "students"
    assert window.state.course_id == course_id
    assert window.pages["students"].proxy.rowCount() == 12


def test_language_preference_persists_after_restart(window, ctx, qapp):
    from attendance_manager.paths import AppPaths
    from attendance_manager.services.context import AppContext
    from attendance_manager.ui.language import apply_language, load_language
    from attendance_manager.ui.main_window import MainWindow

    window.change_language("ar")
    root = ctx.paths.root
    restarted = AppContext(AppPaths(root))  # a fresh process would do exactly this
    try:
        assert load_language(restarted) == "ar"
        apply_language(qapp, load_language(restarted))
        second = MainWindow(restarted)
        assert second.isRightToLeft()
        assert second.nav_buttons["students"].text().strip() == "الطلاب"
        second.close()
    finally:
        restarted.close()


def test_arabic_report_keeps_data_intact(ctx, tmp_path):
    course = create_course(ctx, SYNTHETIC / "00.xls")
    add_session(ctx, course.id, SYNTHETIC / "11.xls")
    set_language("ar")
    path = ctx.reports.generate(course.id, tmp_path / "ar.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == [tr(f"report.sheet.{i}") for i in SHEET_IDS]
    ws = wb[tr("report.sheet.cumulative")]
    assert ws.sheet_view.rightToLeft
    headers = [c.value for c in ws[1]]
    assert headers[0] == "الرقم الجامعي" and "Major" in headers  # master-file columns untouched
    first = next(ws.iter_rows(min_row=2, values_only=True))
    assert first[0] == "202400001" and first[1] == "ADAM SAMI KHALIL"  # IDs and names untouched
    grid = wb[tr("report.sheet.by_session")]
    row = next(r for r in grid.iter_rows(min_row=2, values_only=True) if r[0] == "202400012")
    assert row[2] == "غ"


def test_arabic_error_messages(tmp_path):
    from attendance_manager.domain.errors import ImportFileError
    from attendance_manager.importing.reader import read_workbook

    set_language("ar")
    empty = tmp_path / "x.xls"
    empty.write_bytes(b"")
    with pytest.raises(ImportFileError) as err:
        read_workbook(empty)
    assert "فارغ" in err.value.message and "x.xls" in err.value.message


def test_bidi_keeps_ltr_data_intact_in_arabic():
    from attendance_manager.i18n import LRM, bidi_value, paragraph
    from attendance_manager.services.report_service import safe_filename

    assert bidi_value("11.xls") == "11.xls"  # English UI: untouched
    set_language("ar")
    assert bidi_value("11.xls") == f"{LRM}11.xls{LRM}"  # isolated inside RTL text
    assert bidi_value("مهارات حياتية") == "مهارات حياتية"  # Arabic data untouched
    assert bidi_value("202400123") == "202400123"  # IDs untouched
    message = tr("errors.import.empty", file="11.xls")
    assert f"{LRM}11.xls{LRM}" in message and strip_bidi(message).count("11.xls") == 1
    assert paragraph("Course").startswith("\u200f")
    name = safe_filename(tr("report.filename", course="Intro · شعبة 1", date="2026-07-18"))
    assert LRM not in name and "\u200f" not in name


def test_clicking_the_language_switch_rebuilds_safely(window, qtbot, qapp):
    """Real mouse clicks on the header switch (the rebuild replaces the clicked button)."""
    from PySide6.QtCore import Qt

    qtbot.mouseClick(window.language_switch.buttons["ar"], Qt.MouseButton.LeftButton)
    qtbot.waitUntil(lambda: window.isRightToLeft(), timeout=5000)
    assert window.windowTitle() == "مدير الحضور"
    qtbot.mouseClick(window.language_switch.buttons["en"], Qt.MouseButton.LeftButton)
    qtbot.waitUntil(lambda: not window.isRightToLeft(), timeout=5000)
    assert window.windowTitle() == "Attendance Manager"
    # and from the Settings page switch as well
    window.navigate("settings")
    from attendance_manager.ui.widgets.language_switch import LanguageSwitch

    switch = window.pages["settings"].findChild(LanguageSwitch)
    qtbot.mouseClick(switch.buttons["ar"], Qt.MouseButton.LeftButton)
    qtbot.waitUntil(lambda: window.isRightToLeft(), timeout=5000)
    assert window.current_page_key() == "settings"
