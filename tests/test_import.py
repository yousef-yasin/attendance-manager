"""Excel import: reading, layout detection and parsing (real + synthetic files)."""

import zipfile
from pathlib import Path

import openpyxl
import pytest

from attendance_manager.domain.errors import ImportFileError
from attendance_manager.importing.detection import ColumnMapping, Confidence
from attendance_manager.importing.importer import (
    build_roster,
    detect_session_date,
    load_file,
    load_roster,
    parse_datetime_text,
    parse_table,
)
from attendance_manager.importing.reader import read_workbook

from .conftest import REAL, SYNTHETIC, requires_real


def parse(path: Path):
    loaded = load_file(path)
    assert loaded.detection.mapping is not None
    return loaded, parse_table(loaded.workbook, loaded.detection.mapping)


# ------------------------------------------------------------------ real files
@requires_real
def test_real_master_structure(real_master):
    loaded, table = parse(real_master)
    assert loaded.workbook.file_format == "xls"
    assert loaded.detection.confidence is Confidence.HIGH
    assert table.sheet_name == "studentsTable"
    assert table.column_names == [
        "#",
        "Student Id",
        "Student Name",
        "Major",
        "Status",
        "Register Date",
    ]
    assert table.id_column_name == "Student Id"
    assert table.name_column_name == "Student Name"
    roster = build_roster(table)
    assert len(roster.students) == 45
    assert roster.info_columns == ["Major", "Status", "Register Date"]
    assert not roster.duplicates
    # Real names/IDs are personal data, so check their shape rather than their values.
    assert all(s.student_number.isdigit() and len(s.student_number) == 9 for s in roster.students)
    assert len({s.student_number for s in roster.students}) == 45
    assert all(s.full_name and s.info["Major"].endswith("/ Bachelor") for s in roster.students)
    # whitespace inside names is normalised (the real file has a name with repeated spaces)
    assert all(s.full_name == " ".join(s.full_name.split()) for s in roster.students)


@requires_real
@pytest.mark.parametrize(
    ("name", "rows", "date"),
    [
        ("11.xls", 29, "2026-07-18"),
        ("22.xls", 29, "2026-07-19"),
        ("33.xls", 37, "2026-07-25"),
        ("44.xls", 38, "2026-07-26"),
    ],
)
def test_real_weekly_teams_reports(name, rows, date):
    loaded, table = parse(REAL / name)
    assert loaded.detection.confidence is Confidence.HIGH
    assert table.id_column_name in ("Email", "Participant ID (UPN)")
    assert table.name_column_name == "Name"
    assert len(table.rows) == rows
    assert str(detect_session_date(table)) == date
    assert table.metadata["Meeting title"].startswith("Meeting with")
    organizer = [r for r in table.rows if not r.id_like]
    # the organizer is identified by an e-mail user name, not a student ID
    assert len(organizer) == 1 and not organizer[0].key.isdigit()


@requires_real
def test_real_33_has_no_header_before_data_but_is_parsed():
    loaded, table = parse(REAL / "33.xls")
    assert loaded.detection.mapping.header_row is None
    # The header row was sorted into the middle of the data: it is skipped, and used for names.
    assert [s.reason for s in table.skipped] == ["Repeated header row"]
    assert table.column_names[4] == "Email"
    assert all(r.key != "email" for r in table.rows)


# ------------------------------------------------------------- synthetic files
def test_synthetic_master():
    roster = load_roster(SYNTHETIC / "00.xls")
    assert len(roster.students) == 12
    keys = [s.key for s in roster.students]
    assert "202400003" in keys  # the padded " 202400003 " cell
    assert roster.students[2].student_number == "202400003"


def test_teams_report_with_header():
    loaded, table = parse(SYNTHETIC / "11.xls")
    assert loaded.detection.mapping.header_row == 9
    assert table.metadata["Start time"] == "7/18/26, 3:56:40 PM"
    assert len(table.rows) == 11  # organizer + 8 students + guest + rejoin


def test_teams_report_without_header_and_sorted_header():
    loaded, table = parse(SYNTHETIC / "22.xls")
    assert loaded.detection.confidence is Confidence.HIGH
    assert loaded.detection.mapping.header_row is None
    assert table.column_names[0] == "Name"
    assert len(table.rows) == 10
    assert [s.reason for s in table.skipped] == ["Repeated header row"]


def test_xlsx_numeric_ids_and_footer():
    loaded, table = parse(SYNTHETIC / "33.xlsx")
    assert loaded.workbook.file_format == "xlsx"
    assert {r.key for r in table.rows} == {f"2024000{i:02d}" for i in range(1, 11)}
    assert any("Total" in s.reason for s in table.skipped)


def test_utf16_tab_separated_csv():
    loaded, table = parse(SYNTHETIC / "44.csv")
    assert loaded.workbook.file_format == "csv"
    assert len(table.rows) == 7
    assert str(detect_session_date(table)) == "2026-07-26"


def test_xls_extension_that_is_really_xlsx(tmp_path):
    target = tmp_path / "renamed.xls"
    target.write_bytes((SYNTHETIC / "33.xlsx").read_bytes())
    loaded, table = parse(target)
    assert loaded.workbook.file_format == "xlsx"
    assert len(table.rows) == 10


def test_html_disguised_as_xls(tmp_path):
    target = tmp_path / "export.xls"
    target.write_text(
        "<html><body><table><tr><th>Student ID</th><th>Name</th></tr>"
        "<tr><td>202400001</td><td>Adam</td></tr><tr><td>202400002</td><td>Basma</td></tr>"
        "<tr><td>202400003</td><td>Carim</td></tr></table></body></html>",
        encoding="utf-8",
    )
    loaded, table = parse(target)
    assert loaded.workbook.file_format == "html"
    assert [r.key for r in table.rows] == ["202400001", "202400002", "202400003"]


def test_duplicate_ids_empty_rows_and_missing_values_in_master(tmp_path):
    path = tmp_path / "master.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["No", "University ID", "Full Name", "Phone"])
    ws.append([1, "202400001", "Adam", "0790000001"])
    ws.append([None, None, None, None])
    ws.append([2, "202400002", None, "0790000002"])
    ws.append([3, " 202400001 ", "Adam again", "0790000003"])
    ws.append([4, "N/A", "Nobody", "0790000004"])
    ws.append([5, 202400003, "Carim", None])
    wb.save(path)
    roster = load_roster(path)
    assert [s.key for s in roster.students] == ["202400001", "202400002", "202400003"]
    assert len(roster.duplicates) == 1
    assert roster.students[1].full_name == "(no name)"
    assert "No" not in roster.students[0].info  # serial column dropped
    assert roster.students[0].info == {"Phone": "0790000001"}
    text = "\n".join(roster.warnings)
    assert "appears more than once" in text
    assert "no name" in text
    assert "does not look like a student ID" in text


def test_ambiguous_columns_need_confirmation(tmp_path):
    path = tmp_path / "ambiguous.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(10):
        ws.append([f"20240{i:04d}", f"07900{i:05d}", f"Person {i} Name"])
    wb.save(path)
    loaded = load_file(path)
    assert loaded.detection.confidence is Confidence.MEDIUM
    assert loaded.needs_confirmation


def test_file_without_ids_is_low_confidence(tmp_path):
    path = tmp_path / "names.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "City"])
    ws.append(["Adam Sami", "Amman"])
    ws.append(["Basma Fuad", "Irbid"])
    wb.save(path)
    loaded = load_file(path)
    assert loaded.detection.confidence is Confidence.LOW
    assert loaded.detection.mapping is None


def test_manual_mapping_overrides_detection():
    loaded = load_file(SYNTHETIC / "11.xls")
    mapping = ColumnMapping(
        sheet_index=0, header_row=9, first_data_row=10, id_column=5, name_column=0
    )
    table = parse_table(loaded.workbook, mapping)
    assert table.id_column_name == "Participant ID (UPN)"
    assert len([r for r in table.rows if r.id_like]) == 10


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"", "empty"),
        (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600, "not a valid Excel"),
        (b"PK\x03\x04garbage", "damaged"),
    ],
)
def test_malformed_files_raise_clear_errors(tmp_path, content, message):
    path = tmp_path / "bad.xls"
    path.write_bytes(content)
    with pytest.raises(ImportFileError) as err:
        read_workbook(path)
    assert message in err.value.message


def test_missing_file_error(tmp_path):
    with pytest.raises(ImportFileError, match="could not be found"):
        read_workbook(tmp_path / "nope.xls")


def test_roster_without_students_is_rejected(tmp_path):
    path = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Student ID", "Name"])
    wb.active.append(["Total", "0"])
    wb.save(path)
    loaded = load_file(path)
    mapping = ColumnMapping(0, 0, 1, 0, 1)
    with pytest.raises(ImportFileError, match="No students"):
        build_roster(parse_table(loaded.workbook, mapping))


def test_zip_that_is_not_a_workbook(tmp_path):
    path = tmp_path / "fake.xlsx"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("hello.txt", "hi")
    with pytest.raises(ImportFileError):
        read_workbook(path)


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("7/18/26, 3:56:40 PM", "2026-07-18"),
        ("7/18/2026, 3:56:40 PM", "2026-07-18"),
        ("2026-07-18 15:56:40", "2026-07-18"),
        ("2026-07-18", "2026-07-18"),
        ("18.07.2026", "2026-07-18"),
    ],
)
def test_parse_datetime_text(text, expected):
    parsed = parse_datetime_text(text)
    assert parsed is not None and parsed.date().isoformat() == expected
