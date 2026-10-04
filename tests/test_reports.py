import openpyxl
import pytest

from attendance_manager.reporting.excel_report import SHEET_NAMES

from .conftest import REAL, SYNTHETIC, add_session, create_course, requires_real


def build(ctx, master, weekly):
    course = create_course(ctx, master)
    for path in weekly:
        add_session(ctx, course.id, path)
    return course


def test_report_has_all_sheets_and_correct_values(ctx, tmp_path):
    course = build(
        ctx,
        SYNTHETIC / "00.xls",
        [SYNTHETIC / n for n in ("11.xls", "22.xls", "33.xlsx", "44.csv")],
    )
    path = ctx.reports.generate(course.id, tmp_path / "report.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == list(SHEET_NAMES)

    cumulative = wb["Cumulative Attendance"]
    headers = [c.value for c in cumulative[1]]
    assert headers[:2] == ["Student ID", "Student Name"]
    assert headers[-5:] == ["Sessions", "Present", "Absent", "Attendance %", "Status"]
    assert "Major" in headers
    rows = {r[0]: r for r in cumulative.iter_rows(min_row=2, values_only=True)}
    assert len(rows) == 12
    s1 = rows["202400001"]
    assert s1[-5:] == (4, 4, 0, 1.0, "Perfect")
    assert rows["202400012"][-5:] == (4, 0, 4, 0.0, "At risk")

    grid = wb["Session by Session"]
    row1 = next(r for r in grid.iter_rows(min_row=2, values_only=True) if r[0] == "202400012")
    assert row1[2:6] == ("A", "A", "A", "A")

    ranking = [r for r in wb["Absence Ranking"].iter_rows(min_row=2, values_only=True)]
    assert ranking[0][1] == "202400012" and ranking[0][0] == 1
    highest = [r for r in wb["Highest Attendance"].iter_rows(min_row=2, values_only=True)]
    assert highest[0][3] == 1.0 and highest[0][0] == 1
    assert sum(1 for r in highest if r[0] == 1) == 4  # ties share rank 1

    at_risk_ids = [
        r[0] for r in wb["At-Risk Students"].iter_rows(min_row=5, values_only=True) if r[0]
    ]
    stats = ctx.statistics.course_statistics(course.id).dashboard
    expected = [s.student.student_number for s in stats.at_risk_students + stats.warning_students]
    assert at_risk_ids == expected

    details = list(wb["Student Details"].iter_rows(min_row=2, values_only=True))
    assert len(details) == 12 * 4
    summary_values = [c.value for c in wb["Summary"]["A"]]
    assert "Attendance Report" in summary_values
    assert wb["Summary"]._charts  # trend + distribution charts present


@requires_real
def test_real_report(ctx, real_master, tmp_path):
    course = build(ctx, real_master, [REAL / n for n in ("11.xls", "22.xls", "33.xls", "44.xls")])
    path = ctx.reports.generate(course.id, tmp_path / "real.xlsx")
    wb = openpyxl.load_workbook(path)
    rows = list(wb["Cumulative Attendance"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 45
    assert sum(r[-4] for r in rows) == 127  # present
    assert sum(r[-3] for r in rows) == 53  # absent


def test_session_export(ctx, tmp_path):
    course = build(ctx, SYNTHETIC / "00.xls", [SYNTHETIC / "11.xls"])
    session = ctx.attendance.sessions(course.id)[0]
    path = ctx.reports.export_session(session.id, tmp_path / "s.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Session"]
    statuses = [r[2] for r in ws.iter_rows(min_row=10, values_only=True) if r[2]]
    assert statuses.count("Present") == 8 and statuses.count("Absent") == 4
    assert "Not on master list" in wb.sheetnames


def test_report_for_course_without_sessions(ctx, tmp_path):
    course = create_course(ctx, SYNTHETIC / "00.xls")
    path = ctx.reports.generate(course.id, tmp_path / "empty.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == list(SHEET_NAMES)


def test_report_locked_file_gives_friendly_error(ctx, tmp_path, monkeypatch):
    from attendance_manager.domain.errors import AttendanceError
    from attendance_manager.reporting import excel_report

    course = create_course(ctx, SYNTHETIC / "00.xls")

    def locked(*_a, **_k):
        raise PermissionError("locked")

    monkeypatch.setattr("attendance_manager.services.report_service.write_course_report", locked)
    with pytest.raises(AttendanceError, match="close it"):
        ctx.reports.generate(course.id, tmp_path / "r.xlsx")
    assert excel_report  # module import sanity
